from __future__ import annotations

import subprocess
import sys
from importlib.machinery import SourceFileLoader
from importlib.util import module_from_spec, spec_from_file_location
from pathlib import Path
from types import ModuleType

from docx import Document

SCRIPT = Path(__file__).resolve().parents[1] / "scripts" / "docmd"


def load_docmd() -> ModuleType:
    loader = SourceFileLoader("docmd_cli", str(SCRIPT))
    spec = spec_from_file_location("docmd_cli", SCRIPT, loader=loader)
    assert spec is not None
    assert spec.loader is not None
    module = module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def test_code_fences(tmp_path: Path) -> None:
    docmd = load_docmd()

    assert docmd.verify_markdown("```\ncode\n```\n") == []
    assert docmd.verify_markdown("```\ncode\n") == ["unbalanced code fence"]

    document = Document()
    document.add_paragraph("```")
    source = tmp_path / "fence.docx"
    document.save(source)
    result = subprocess.run(
        [sys.executable, str(SCRIPT), str(source), "--stdout", "--verify"],
        capture_output=True,
        text=True,
    )
    assert result.returncode != 0
    assert "unbalanced code fence" in result.stderr


def test_table_alignment() -> None:
    docmd = load_docmd()

    valid = "| A | B |\n| --- | --- |\n| C \\| D | E |\n"
    assert docmd.verify_markdown(valid) == []

    missing = "| A | B |\n| --- | --- |\n| C |\n"
    assert docmd.verify_markdown(missing) == [
        "table row column count does not match header"
    ]

    extra = "| A | B |\n| --- | --- |\n| C | D | E |\n"
    assert docmd.verify_markdown(extra) == [
        "table row column count does not match header"
    ]

    fenced = "```\n| A | B |\n| --- | --- |\n| C |\n```\n"
    assert docmd.verify_markdown(fenced) == []


def test_asset_links(tmp_path: Path) -> None:
    docmd = load_docmd()

    assert docmd.verify_markdown("![x](missing.png)", tmp_path) == [
        "missing asset: missing.png"
    ]
    absolute = tmp_path / "absolute-missing.png"
    assert docmd.verify_markdown(f"![x]({absolute})", tmp_path) == [
        f"missing asset: {absolute}"
    ]
    assert docmd.verify_markdown("[x](missing.png)", tmp_path) == [
        "missing asset: missing.png"
    ]

    empty = tmp_path / "empty.png"
    empty.write_bytes(b"")
    assert docmd.verify_markdown("![x](empty.png)", tmp_path) == [
        "zero-byte asset: empty.png"
    ]

    valid = tmp_path / "valid.png"
    valid.write_bytes(b"image")
    assert docmd.verify_markdown("![x](valid.png)", tmp_path) == []

    assert docmd.verify_markdown("![x](https://example.com/a.png)", tmp_path) == []
    assert docmd.verify_markdown("![x](data:image/png;base64,AAAA)", tmp_path) == []
    assert docmd.verify_markdown("![x](#anchor)", tmp_path) == []
    assert docmd.verify_markdown("```\n![x](missing.png)\n```", tmp_path) == []


def test_link_targets(tmp_path: Path) -> None:
    docmd = load_docmd()

    for target in (
        "#anchor",
        "https://example.com/a.png",
        "http://example.com/a.png",
        "mailto:someone@example.com",
        "data:image/png;base64,AAAA",
        "file:///tmp/a.png",
        "//example.com/a.png",
    ):
        assert docmd.verify_markdown(f"![x]({target})", tmp_path) == []

    valid = tmp_path / "valid.png"
    valid.write_bytes(b"image")
    assert docmd.verify_markdown("![x](valid.png)", tmp_path) == []
    assert docmd.verify_markdown("![x](valid.png#section)", tmp_path) == []
    assert docmd.verify_markdown("![x](valid.png?raw=1)", tmp_path) == []

    (tmp_path / "my image.png").write_bytes(b"image")
    assert docmd.verify_markdown("![x](my%20image.png)", tmp_path) == []
    assert docmd.verify_markdown("![x](missing%20image.png)", tmp_path) == [
        "missing asset: missing%20image.png"
    ]


def test_tilde_fences(tmp_path: Path) -> None:
    docmd = load_docmd()

    assert docmd.verify_markdown("~~~\ncode\n~~~\n") == []
    assert docmd.verify_markdown("~~~\ncode\n") == ["unbalanced code fence"]
    assert docmd.verify_markdown("~~~\n![x](missing.png)\n~~~\n") == []
    assert docmd.verify_markdown("~~~\n| A | B |\n| --- | --- |\n| C |\n~~~\n") == []
    assert docmd.verify_markdown("```\n~~~\n```\n") == []

    document = Document()
    document.add_paragraph("~~~")
    source = tmp_path / "tilde.docx"
    document.save(source)
    result = subprocess.run(
        [sys.executable, str(SCRIPT), str(source), "--stdout", "--verify"],
        capture_output=True,
        text=True,
    )
    assert result.returncode != 0
    assert "unbalanced code fence" in result.stderr
