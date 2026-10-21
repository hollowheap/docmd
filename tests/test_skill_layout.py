from __future__ import annotations

import os
from pathlib import Path

import pymupdf
from docx import Document
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

ROOT = Path(__file__).resolve().parents[1]
SKILL_ROOT = ROOT / "skills" / "docmd"


def test_skill_and_readme() -> None:
    skill = (SKILL_ROOT / "SKILL.md").read_text(encoding="utf-8")
    frontmatter = skill.removeprefix("---\n").split("\n---", 1)[0]

    assert "name: docmd" in frontmatter
    assert "description:" in frontmatter
    readme = (SKILL_ROOT / "scripts" / "README.md").read_text(encoding="utf-8")
    assert readme.strip()
    assert (SKILL_ROOT / "README.md").read_text(encoding="utf-8").strip()


def test_reference_assets() -> None:
    reference = SKILL_ROOT / "references"

    for name in ("docx.md", "pptx.md", "pdf.md"):
        guide = reference / name
        assert guide.is_file()
        assert guide.read_text(encoding="utf-8").strip()

    for name in ("sample.docx", "sample.pptx", "sample.pdf"):
        fixture = reference / "fixtures" / name
        assert fixture.is_file()
        assert fixture.stat().st_size > 0


def test_executable_script() -> None:
    packaged = SKILL_ROOT / "scripts" / "docmd"

    assert packaged.is_file()
    assert os.access(packaged, os.X_OK)
    assert packaged.read_text(encoding="utf-8") == (
        ROOT / "scripts" / "docmd"
    ).read_text(encoding="utf-8")

    script = packaged.read_text(encoding="utf-8")
    assert script.startswith("#!/usr/bin/env -S uv run --script")
    assert "# /// script" in script
    assert 'requires-python = ">=3.11"' in script
    assert "dependencies = [" in script
    assert script.count("# ///") >= 2


def test_frontmatter_contract() -> None:
    skill = (SKILL_ROOT / "SKILL.md").read_text(encoding="utf-8")
    frontmatter = skill.removeprefix("---\n").split("\n---", 1)[0]

    assert "name: docmd" in frontmatter
    assert SKILL_ROOT.name == "docmd"
    assert "license: MIT" in frontmatter
    assert "compatibility:" in frontmatter
    assert "version:" in frontmatter

    description = next(
        line.split(":", 1)[1].strip()
        for line in frontmatter.splitlines()
        if line.startswith("description:")
    )
    assert description
    assert len(description) <= 1024

    for guide in ("references/docx.md", "references/pptx.md", "references/pdf.md"):
        assert guide in skill


def test_fixture_validity() -> None:
    fixtures = SKILL_ROOT / "references" / "fixtures"

    document = Document(fixtures / "sample.docx")
    assert any(paragraph.text.strip() for paragraph in document.paragraphs)
    assert len(document.part.package.image_parts) >= 1

    presentation = Presentation(fixtures / "sample.pptx")
    assert presentation.slides
    assert any(
        shape.has_text_frame and shape.text_frame.text.strip()
        for slide in presentation.slides
        for shape in slide.shapes
    )
    assert any(
        shape.shape_type == MSO_SHAPE_TYPE.PICTURE
        for slide in presentation.slides
        for shape in slide.shapes
    )

    pdf = pymupdf.open(fixtures / "sample.pdf")
    assert pdf.page_count >= 1
    assert pdf[0].get_text().strip()
    assert pdf[0].get_images()
    pdf.close()
