from __future__ import annotations

import base64
import subprocess
import sys
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.shared import Pt

SCRIPT = Path(__file__).resolve().parents[1] / "scripts" / "docmd"
PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAusB9Y9Zl1sAAAAASUVORK5CYII="
)


def run_docmd(*args: str) -> subprocess.CompletedProcess[str]:
    return subprocess.run(
        [sys.executable, str(SCRIPT), *args],
        capture_output=True,
        text=True,
    )


def test_block_extraction(tmp_path: Path) -> None:
    document = Document()
    document.add_heading("Title", level=1)
    document.add_paragraph("Body text")
    document.add_paragraph("First item", style="List Bullet")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Name"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "A"
    table.cell(1, 1).text = "B"
    document.add_picture(BytesIO(PNG))
    source = tmp_path / "sample.docx"
    document.save(source)

    result = run_docmd(str(source), "--stdout")

    assert result.returncode == 0, result.stderr
    assert "# Title" in result.stdout
    assert "Body text" in result.stdout
    assert "- First item" in result.stdout
    assert "| Name | Value |\n| --- | --- |\n| A | B |" in result.stdout
    assert "data:image/png;base64," in result.stdout
    assert result.stdout.index("# Title") < result.stdout.index("Body text")
    assert result.stdout.index("Body text") < result.stdout.index("- First item")


def test_heading_promotion(tmp_path: Path) -> None:
    document = Document()
    for level in range(1, 7):
        document.add_heading(f"Heading level {level}", level=level)

    document.add_paragraph("Body text")
    for text, size in (
        ("Large bold", 20),
        ("Medium bold", 16),
        ("Same size bold", 11),
    ):
        paragraph = document.add_paragraph()
        run = paragraph.add_run(text)
        run.bold = True
        run.font.size = Pt(size)

    source = tmp_path / "headings.docx"
    document.save(source)

    result = run_docmd(str(source), "--stdout")

    assert result.returncode == 0, result.stderr
    for level in range(1, 7):
        assert f"{'#' * level} Heading level {level}" in result.stdout
    assert "# Large bold" in result.stdout
    assert "## Medium bold" in result.stdout
    assert "Same size bold" in result.stdout
    assert "## Same size bold" not in result.stdout


def test_numbered_list_markers(tmp_path: Path) -> None:
    document = Document()
    document.add_paragraph("First bullet", style="List Bullet")
    document.add_paragraph("First number", style="List Number")
    document.add_paragraph("Second number", style="List Number")
    source = tmp_path / "lists.docx"
    document.save(source)

    result = run_docmd(str(source), "--stdout")

    assert result.returncode == 0, result.stderr
    assert "- First bullet" in result.stdout
    assert "1. First number" in result.stdout
    assert "1. Second number" in result.stdout
    assert (
        result.stdout.index("- First bullet")
        < result.stdout.index("1. First number")
        < result.stdout.index("1. Second number")
    )


def test_image_only_paragraph(tmp_path: Path) -> None:
    document = Document()
    document.add_paragraph("Before")
    document.add_picture(BytesIO(PNG))
    source = tmp_path / "image.docx"
    document.save(source)

    result = run_docmd(str(source), "--stdout")

    assert result.returncode == 0, result.stderr
    assert "\n\n\n" not in result.stdout
    blocks = result.stdout.strip().split("\n\n")
    assert blocks[0] == "Before"
    assert len(blocks) == 2
    assert blocks[1].startswith("![](data:image/png;base64,")
