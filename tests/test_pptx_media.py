from __future__ import annotations

import base64
import subprocess
import sys
from importlib.machinery import SourceFileLoader
from importlib.util import module_from_spec, spec_from_file_location
from io import BytesIO
from pathlib import Path
from types import ModuleType

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches

SCRIPT = Path(__file__).resolve().parents[1] / "scripts" / "docmd"
PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAusB9Y9Zl1sAAAAASUVORK5CYII="
)
BLOB = b"picture-bytes"
ENCODED = base64.b64encode(BLOB).decode("ascii")


def run_docmd(*args: str) -> subprocess.CompletedProcess[str]:
    return subprocess.run(
        [sys.executable, str(SCRIPT), *args],
        capture_output=True,
        text=True,
    )


def load_docmd() -> ModuleType:
    loader = SourceFileLoader("docmd_cli", str(SCRIPT))
    spec = spec_from_file_location("docmd_cli", SCRIPT, loader=loader)
    assert spec is not None
    assert spec.loader is not None
    module = module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


class FakeImage:
    def __init__(self, content_type: str | None, extension: str | None) -> None:
        self.content_type = content_type
        self.ext = extension
        self.blob = BLOB


class FakePicture:
    shape_type = MSO_SHAPE_TYPE.PICTURE
    has_text_frame = False
    has_table = False

    def __init__(
        self, content_type: str | None, extension: str | None, shape_id: int
    ) -> None:
        self.shape_id = shape_id
        self.image = FakeImage(content_type, extension)


class BrokenPicture:
    shape_type = MSO_SHAPE_TYPE.PICTURE
    has_text_frame = False
    has_table = False
    shape_id = 9

    @property
    def image(self) -> FakeImage:
        raise KeyError("no image")


class FakeShapes(list):
    title = None


class FakeSlide:
    has_notes_slide = False

    def __init__(self, shapes: list[object]) -> None:
        self.shapes = FakeShapes(shapes)


class FakePresentation:
    def __init__(self, shapes: list[object]) -> None:
        self.slides = [FakeSlide(shapes)]


def test_picture_mime(tmp_path: Path, monkeypatch) -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_picture(BytesIO(PNG), Inches(1), Inches(1))
    source = tmp_path / "picture.pptx"
    presentation.save(source)

    result = run_docmd(str(source), "--stdout")

    assert result.returncode == 0, result.stderr
    assert "![](data:image/png;base64," in result.stdout

    docmd = load_docmd()
    shapes = [
        FakePicture("image/png", "png", 1),
        FakePicture(None, "jpeg", 2),
        FakePicture(None, "xyz", 3),
        FakePicture(None, "", 4),
        BrokenPicture(),
    ]
    monkeypatch.setattr(
        docmd,
        "open_office_source",
        lambda source, password, opener: FakePresentation(shapes),
    )

    markdown = docmd.convert_pptx(Path("fake.pptx"), include_images=True)

    assert markdown.count("![](data:image/") == 4
    assert f"data:image/png;base64,{ENCODED}" in markdown
    assert f"data:image/jpeg;base64,{ENCODED}" in markdown
    assert f"data:image/xyz;base64,{ENCODED}" in markdown
    assert markdown.count("data:image/png;base64,") == 2
