from __future__ import annotations

import base64
import re
import subprocess
import sys
from importlib.machinery import SourceFileLoader
from importlib.util import module_from_spec, spec_from_file_location
from io import BytesIO
from pathlib import Path
from types import ModuleType

import pymupdf
from docx import Document
from pptx import Presentation

SCRIPT = Path(__file__).resolve().parents[1] / "scripts" / "docmd"
PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAusB9Y9Zl1sAAAAASUVORK5CYII="
)


def run_docmd(*args: str) -> subprocess.CompletedProcess[str]:
    return subprocess.run(
        [sys.executable, str(SCRIPT), *args],
        capture_output=True,
        text=True,
    )


def test_recursive_mirroring(tmp_path: Path) -> None:
    source = tmp_path / "input"
    nested = source / "nested"
    deep = nested / "deep"
    deep.mkdir(parents=True)
    hidden_dir = source / ".cache"
    hidden_dir.mkdir()

    document = Document()
    document.add_paragraph("Top Doc")
    document.save(source / "top.docx")

    hidden = Document()
    hidden.add_paragraph("Hidden Doc")
    hidden.save(source / ".hidden.docx")

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Nested Slides"
    presentation.save(nested / "slides.pptx")

    pdf = pymupdf.open()
    page = pdf.new_page()
    page.insert_text((72, 72), "PDF Body Text")
    pdf.save(deep / "paper.pdf")
    pdf.close()

    hidden_pdf = pymupdf.open()
    page = hidden_pdf.new_page()
    page.insert_text((72, 72), "Hidden PDF")
    hidden_pdf.save(hidden_dir / "hidden.pdf")
    hidden_pdf.close()

    (source / "notes.txt").write_text("not a document")
    out = tmp_path / "out"

    result = run_docmd(str(source), "-o", str(out), "--no-ocr")

    assert result.returncode == 0, result.stderr
    assert "Top Doc" in (out / "top.md").read_text()
    assert "Nested Slides" in (out / "nested" / "slides.md").read_text()
    assert "PDF Body Text" in (out / "nested" / "deep" / "paper.md").read_text()
    assert not (out / ".hidden.md").exists()
    assert not (out / ".cache").exists()
    assert not (out / "notes.md").exists()

    stdout_result = run_docmd(str(source), "--stdout", "--no-ocr")
    assert stdout_result.returncode == 1
    assert "--stdout is not supported for directory input" in stdout_result.stderr


def test_asset_co_location(tmp_path: Path) -> None:
    nested = tmp_path / "input" / "nested"
    nested.mkdir(parents=True)

    document = Document()
    document.add_paragraph("Before")
    document.add_picture(BytesIO(PNG))
    document.save(nested / "doc.docx")

    out = tmp_path / "out"
    result = run_docmd(str(tmp_path / "input"), "-o", str(out))

    assert result.returncode == 0, result.stderr
    markdown_path = out / "nested" / "doc.md"
    markdown = markdown_path.read_text()
    assert "data:image" not in markdown
    assert not (out / "doc.assets").exists()

    assets = out / "nested" / "doc.assets"
    files = list(assets.iterdir())
    assert len(files) == 1
    assert files[0].read_bytes() == PNG

    link = re.search(r"!\[[^\]]*\]\(([^)]+)\)", markdown)
    assert link is not None
    assert (markdown_path.parent / link.group(1)).resolve() == files[0].resolve()


def test_partial_failure(tmp_path: Path) -> None:
    source = tmp_path / "input"
    source.mkdir()

    good = Document()
    good.add_paragraph("Good Doc")
    good.save(source / "a.docx")

    (source / "b.docx").write_bytes(b"not a valid document")

    out = tmp_path / "out"
    result = run_docmd(str(source), "-o", str(out))

    assert result.returncode == 2
    assert (out / "a.md").is_file()
    assert "Good Doc" in (out / "a.md").read_text(encoding="utf-8")
    assert not (out / "b.md").exists()


def test_parallel_jobs(tmp_path: Path) -> None:
    source = tmp_path / "input"
    nested = source / "nested"
    nested.mkdir(parents=True)
    for name, text in (("a.docx", "Alpha"), ("b.docx", "Beta")):
        document = Document()
        document.add_paragraph(text)
        document.save(source / name)
    document = Document()
    document.add_paragraph("Gamma")
    document.save(nested / "c.docx")
    out = tmp_path / "out"

    result = run_docmd(str(source), "-o", str(out), "--jobs", "2", "--no-ocr")

    assert result.returncode == 0, result.stderr
    assert "Alpha" in (out / "a.md").read_text(encoding="utf-8")
    assert "Beta" in (out / "b.md").read_text(encoding="utf-8")
    assert "Gamma" in (out / "nested" / "c.md").read_text(encoding="utf-8")


def test_parallel_jobs_failure(tmp_path: Path) -> None:
    source = tmp_path / "input"
    source.mkdir()
    good = Document()
    good.add_paragraph("Good Doc")
    good.save(source / "a.docx")
    (source / "b.docx").write_bytes(b"not a valid document")
    out = tmp_path / "out"

    result = run_docmd(str(source), "-o", str(out), "--jobs", "2")

    assert result.returncode == 2
    assert not (out / "b.md").exists()


def test_invalid_jobs() -> None:
    result = run_docmd("missing.docx", "--jobs", "0")

    assert result.returncode == 1
    assert "--jobs must be at least 1" in result.stderr


def load_docmd() -> ModuleType:
    loader = SourceFileLoader("docmd_cli", str(SCRIPT))
    spec = spec_from_file_location("docmd_cli", SCRIPT, loader=loader)
    assert spec is not None
    assert spec.loader is not None
    module = module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def test_empty_input(tmp_path: Path) -> None:
    source = tmp_path / "input"
    (source / "nested").mkdir(parents=True)
    out = tmp_path / "out"

    result = run_docmd(str(source), "-o", str(out))

    assert result.returncode == 0, result.stderr
    assert out.is_dir()
    assert list(out.iterdir()) == []

    document = Document()
    document.add_paragraph("Body")
    document.save(source / "b.docx")
    document.save(source / "a.docx")
    document.save(source / "nested" / "a.docx")

    docmd = load_docmd()
    found = docmd.discover_documents(source)

    assert [path.relative_to(source).as_posix() for path in found] == [
        "a.docx",
        "b.docx",
        "nested/a.docx",
    ]


def test_output_inside_input(tmp_path: Path) -> None:
    source = tmp_path / "input"
    out = source / "out"
    out.mkdir(parents=True)

    first = Document()
    first.add_paragraph("Input Doc")
    first.save(source / "a.docx")

    nested = Document()
    nested.add_paragraph("Nested Doc")
    nested.save(out / "b.docx")

    result = run_docmd(str(source), "-o", str(out), "--no-ocr")

    assert result.returncode == 0, result.stderr
    assert (out / "a.md").is_file()
    assert not (out / "b.md").exists()
    assert not (out / "out").exists()

    repeat = run_docmd(str(source), "-o", str(out), "--no-ocr")

    assert repeat.returncode == 0, repeat.stderr
    assert not (out / "out").exists()


def test_verify_directory(tmp_path: Path) -> None:
    source = tmp_path / "input"
    source.mkdir()

    with_image = Document()
    with_image.add_paragraph("Before")
    with_image.add_picture(BytesIO(PNG))
    with_image.save(source / "a.docx")

    out = tmp_path / "out"
    result = run_docmd(str(source), "-o", str(out), "--verify")

    assert result.returncode == 0, result.stderr
    assert (out / "a.md").is_file()

    broken_source = tmp_path / "broken"
    broken_source.mkdir()
    broken = Document()
    broken.add_paragraph("```")
    broken.save(broken_source / "a.docx")
    normal = Document()
    normal.add_paragraph("Body")
    normal.save(broken_source / "b.docx")

    broken_out = tmp_path / "broken-out"
    broken_result = run_docmd(str(broken_source), "-o", str(broken_out), "--verify")

    assert broken_result.returncode == 1
    assert "unbalanced code fence" in broken_result.stderr
    assert (broken_out / "a.md").is_file()
