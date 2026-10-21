from __future__ import annotations

import base64
import os
import re
import subprocess
import sys
from importlib.machinery import SourceFileLoader
from importlib.util import module_from_spec, spec_from_file_location
from io import BytesIO
from pathlib import Path
from types import ModuleType

from docx import Document

SCRIPT = Path(__file__).resolve().parents[1] / "scripts" / "docmd"
PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAusB9Y9Zl1sAAAAASUVORK5CYII="
)
ENCODED_PNG = base64.b64encode(PNG).decode("ascii")


def load_docmd() -> ModuleType:
    loader = SourceFileLoader("docmd_cli", str(SCRIPT))
    spec = spec_from_file_location("docmd_cli", SCRIPT, loader=loader)
    assert spec is not None
    assert spec.loader is not None
    module = module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def run_docmd(*args: str) -> subprocess.CompletedProcess[str]:
    return subprocess.run(
        [sys.executable, str(SCRIPT), *args],
        capture_output=True,
        text=True,
    )


def test_file_assets(tmp_path: Path) -> None:
    document = Document()
    document.add_paragraph("Before")
    document.add_picture(BytesIO(PNG))
    document.add_paragraph("After")
    source = tmp_path / "sample.docx"
    document.save(source)
    output = tmp_path / "report.md"

    result = run_docmd(str(source), "-o", str(output))

    assert result.returncode == 0, result.stderr
    markdown = output.read_text()
    assert "data:image" not in markdown

    assets = list((tmp_path / "report.assets").iterdir())
    assert len(assets) == 1
    assert assets[0].read_bytes() == PNG

    link = re.search(r"!\[[^\]]*\]\(([^)]+)\)", markdown)
    assert link is not None
    assert (output.parent / link.group(1)).resolve() == assets[0].resolve()


def test_stdout_data_uris(tmp_path: Path) -> None:
    document = Document()
    document.add_picture(BytesIO(PNG))
    source = tmp_path / "sample.docx"
    document.save(source)

    result = run_docmd(str(source), "--stdout")

    assert result.returncode == 0, result.stderr
    match = re.search(r"!\[[^\]]*\]\(data:image/png;base64,([^)]+)\)", result.stdout)
    assert match is not None
    assert base64.b64decode(match.group(1)) == PNG
    assert not (tmp_path / "sample.assets").exists()


def test_no_images(tmp_path: Path) -> None:
    document = Document()
    document.add_paragraph("Before")
    document.add_picture(BytesIO(PNG))
    document.add_paragraph("After")
    source = tmp_path / "sample.docx"
    document.save(source)

    result = run_docmd(str(source), "--stdout", "--no-images")

    assert result.returncode == 0, result.stderr
    assert "![" not in result.stdout
    assert "data:image" not in result.stdout
    assert "Before" in result.stdout
    assert "After" in result.stdout

    output = tmp_path / "report.md"
    file_result = run_docmd(str(source), "-o", str(output), "--no-images")

    assert file_result.returncode == 0, file_result.stderr
    assert "![" not in output.read_text()
    assert not (tmp_path / "report.assets").exists()


def test_asset_link_with_spaces(tmp_path: Path) -> None:
    document = Document()
    document.add_picture(BytesIO(PNG))
    source = tmp_path / "doc.docx"
    document.save(source)
    output = tmp_path / "my report.md"

    result = run_docmd(str(source), "-o", str(output), "--verify")

    assert result.returncode == 0, result.stderr
    assert "](my%20report.assets/image-001.png)" in output.read_text(encoding="utf-8")


def test_utf8_output_under_c_locale(tmp_path: Path) -> None:
    document = Document()
    document.add_paragraph("café 中文")
    source = tmp_path / "unicode.docx"
    document.save(source)
    output = tmp_path / "unicode.md"
    env = dict(
        os.environ,
        LC_ALL="C",
        LANG="C",
        PYTHONUTF8="0",
        PYTHONCOERCECLOCALE="0",
    )

    file_result = subprocess.run(
        [sys.executable, str(SCRIPT), str(source), "-o", str(output)],
        capture_output=True,
        env=env,
    )

    assert file_result.returncode == 0, file_result.stderr
    assert "café 中文" in output.read_text(encoding="utf-8")

    stdout_result = subprocess.run(
        [sys.executable, str(SCRIPT), str(source), "--stdout"],
        capture_output=True,
        env=env,
    )

    assert stdout_result.returncode == 0, stdout_result.stderr
    assert "café 中文" in stdout_result.stdout.decode("utf-8")


def test_asset_naming(tmp_path: Path) -> None:
    document = Document()
    document.add_picture(BytesIO(PNG))
    document.add_picture(BytesIO(PNG))
    source = tmp_path / "sample.docx"
    document.save(source)
    output = tmp_path / "report.md"
    assets = tmp_path / "report.assets"
    assets.mkdir()
    (assets / "keep.txt").write_text("keep", encoding="utf-8")

    result = run_docmd(str(source), "-o", str(output))

    assert result.returncode == 0, result.stderr
    assert sorted(path.name for path in assets.iterdir()) == [
        "image-001.png",
        "image-002.png",
        "keep.txt",
    ]

    docmd = load_docmd()
    unknown = tmp_path / "unknown.md"
    docmd.write_markdown(
        f"![x](data:image/png;base64,{ENCODED_PNG})\n"
        f"![y](data:image/x-icon;base64,{ENCODED_PNG})\n",
        unknown,
    )

    assert sorted(path.name for path in (tmp_path / "unknown.assets").iterdir()) == [
        "image-001.png",
        "image-002.bin",
    ]


def test_non_image_data_uri(tmp_path: Path) -> None:
    docmd = load_docmd()
    output = tmp_path / "report.md"
    markdown = (
        "![x](data:text/plain;base64,QUJD)\n"
        "![y](data:image/png;base64,not*base64)\n"
    )

    docmd.write_markdown(markdown, output)

    assert output.read_text(encoding="utf-8") == markdown
    assert not (tmp_path / "report.assets").exists()
