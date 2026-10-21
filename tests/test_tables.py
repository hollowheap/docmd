from __future__ import annotations

import subprocess
import sys
from pathlib import Path

from docx import Document
from pptx import Presentation
from pptx.util import Inches

SCRIPT = Path(__file__).resolve().parents[1] / "scripts" / "docmd"


def run_docmd(*args: str) -> subprocess.CompletedProcess[str]:
    return subprocess.run(
        [sys.executable, str(SCRIPT), *args],
        capture_output=True,
        text=True,
    )


def assert_normalized(markdown: str) -> None:
    expected = [
        "| Merged | Merged | Header 3 |",
        "| --- | --- | --- |",
        "| Line 1<br>Line 2 | A\\|B | C |",
    ]
    lines = markdown.splitlines()
    start = lines.index(expected[0])
    assert lines[start : start + len(expected)] == expected


def test_docx_table_normalization(tmp_path: Path) -> None:
    document = Document()
    table = document.add_table(rows=2, cols=3)
    table.cell(0, 0).merge(table.cell(0, 1))
    table.cell(0, 0).text = "Merged"
    table.cell(0, 2).text = "Header 3"
    table.cell(1, 0).text = "Line 1\nLine 2"
    table.cell(1, 1).text = "A|B"
    table.cell(1, 2).text = "C"
    source = tmp_path / "table.docx"
    document.save(source)

    result = run_docmd(str(source), "--stdout")

    assert result.returncode == 0, result.stderr
    assert_normalized(result.stdout)


def test_pptx_table_normalization(tmp_path: Path) -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    table = slide.shapes.add_table(2, 3, Inches(1), Inches(1), Inches(4), Inches(1)).table
    table.cell(0, 0).merge(table.cell(0, 1))
    table.cell(0, 0).text = "Merged"
    table.cell(0, 2).text = "Header 3"
    table.cell(1, 0).text = "Line 1\nLine 2"
    table.cell(1, 1).text = "A|B"
    table.cell(1, 2).text = "C"
    source = tmp_path / "table.pptx"
    presentation.save(source)

    result = run_docmd(str(source), "--stdout")

    assert result.returncode == 0, result.stderr
    assert_normalized(result.stdout)


def test_empty_tables(tmp_path: Path) -> None:
    document = Document()
    document.add_paragraph("Before")
    document.add_table(rows=0, cols=0)
    document.add_paragraph("After")
    source = tmp_path / "empty.docx"
    document.save(source)

    result = run_docmd(str(source), "--stdout")

    assert result.returncode == 0, result.stderr
    assert "|" not in result.stdout
    assert "Before" in result.stdout
    assert "After" in result.stdout

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_table(1, 1, Inches(1), Inches(1), Inches(1), Inches(1))
    presentation_source = tmp_path / "empty.pptx"
    presentation.save(presentation_source)

    pptx_result = run_docmd(str(presentation_source), "--stdout")

    assert pptx_result.returncode == 0, pptx_result.stderr
    assert "|  |\n| --- |" in pptx_result.stdout
