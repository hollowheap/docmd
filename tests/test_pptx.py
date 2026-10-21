from __future__ import annotations

import subprocess
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

SCRIPT = Path(__file__).resolve().parents[1] / "scripts" / "docmd"


def run_docmd(*args: str) -> subprocess.CompletedProcess[str]:
    return subprocess.run(
        [sys.executable, str(SCRIPT), *args],
        capture_output=True,
        text=True,
    )


def test_slide_extraction(tmp_path: Path) -> None:
    presentation = Presentation()
    first = presentation.slides.add_slide(presentation.slide_layouts[1])
    first.shapes.title.text = "Intro"
    first.placeholders[1].text = "First point"
    first.notes_slide.notes_text_frame.text = "Remember this"

    second = presentation.slides.add_slide(presentation.slide_layouts[5])
    second.shapes.title.text = "Data"
    table = second.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(4), Inches(2)).table
    table.cell(0, 0).text = "Name"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "A"
    table.cell(1, 1).text = "B"

    source = tmp_path / "sample.pptx"
    presentation.save(source)

    result = run_docmd(str(source), "--stdout")

    assert result.returncode == 0, result.stderr
    assert "Intro" in result.stdout
    assert "First point" in result.stdout
    assert "Remember this" in result.stdout
    assert "Data" in result.stdout
    assert "| Name | Value |\n| --- | --- |\n| A | B |" in result.stdout
    assert result.stdout.index("Intro") < result.stdout.index("Data")


def test_markdown_layout(tmp_path: Path) -> None:
    presentation = Presentation()
    first = presentation.slides.add_slide(presentation.slide_layouts[1])
    first.shapes.title.text = "Intro"
    first.placeholders[1].text = "First point"
    first.notes_slide.notes_text_frame.text = "Remember this"

    second = presentation.slides.add_slide(presentation.slide_layouts[1])
    second.shapes.title.text = "Data"
    second.placeholders[1].text = "Second point"

    source = tmp_path / "layout.pptx"
    presentation.save(source)

    result = run_docmd(str(source), "--stdout")

    assert result.returncode == 0, result.stderr
    assert "## Slide 1: Intro" in result.stdout
    assert "## Slide 2: Data" in result.stdout
    assert "> **Speaker Notes:** Remember this" in result.stdout
    assert (
        result.stdout.index("## Slide 1: Intro")
        < result.stdout.index("---")
        < result.stdout.index("## Slide 2: Data")
    )
