from __future__ import annotations

import os
from pathlib import Path

import pymupdf
from docx import Document
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

ROOT = Path(__file__).resolve().parents[1]


def test_skill_and_readme() -> None:
    skill = (ROOT / "dist" / "SKILL.md").read_text()
    frontmatter = skill.removeprefix("---\n").split("\n---", 1)[0]

    assert "name: docmd" in frontmatter
    assert "description:" in frontmatter
    readme = (ROOT / "dist" / "scripts" / "README.md").read_text()
    assert readme.strip()


def test_reference_assets() -> None:
    reference = ROOT / "dist" / "reference"

    for name in ("docx.md", "pptx.md", "pdf.md"):
        guide = reference / name
        assert guide.is_file()
        assert guide.read_text().strip()

    for name in ("sample.docx", "sample.pptx", "sample.pdf"):
        fixture = reference / "fixtures" / name
        assert fixture.is_file()
        assert fixture.stat().st_size > 0


def test_executable_script() -> None:
    packaged = ROOT / "dist" / "scripts" / "docmd"

    assert packaged.is_file()
    assert os.access(packaged, os.X_OK)
    assert packaged.read_text() == (ROOT / "scripts" / "docmd").read_text()

    script = packaged.read_text()
    assert script.startswith("#!/usr/bin/env -S uv run --script")
    assert "# /// script" in script
    assert 'requires-python = ">=3.11"' in script
    assert "dependencies = [" in script
    assert script.count("# ///") >= 2


def test_fixture_validity() -> None:
    fixtures = ROOT / "dist" / "reference" / "fixtures"

    document = Document(fixtures / "sample.docx")
    assert any(paragraph.text.strip() for paragraph in document.paragraphs)
    assert len(document.part.package.image_parts) >= 1

    presentation = Presentation(fixtures / "sample.pptx")
    assert presentation.slides
    assert any(
        shape.has_text_frame and shape.text_frame.text.strip()
        for slide in presentation.slides
        for shape in slide.shapes
    )
    assert any(
        shape.shape_type == MSO_SHAPE_TYPE.PICTURE
        for slide in presentation.slides
        for shape in slide.shapes
    )

    pdf = pymupdf.open(fixtures / "sample.pdf")
    assert pdf.page_count >= 1
    assert pdf[0].get_text().strip()
    assert pdf[0].get_images()
    pdf.close()
