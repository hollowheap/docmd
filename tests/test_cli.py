from __future__ import annotations

import subprocess
import sys
from pathlib import Path

from docx import Document

SCRIPT = Path(__file__).resolve().parents[1] / "scripts" / "docmd"


def run_docmd(*args: str) -> subprocess.CompletedProcess[str]:
    return subprocess.run(
        [sys.executable, str(SCRIPT), *args],
        capture_output=True,
        text=True,
    )


def test_option_contract() -> None:
    result = run_docmd("--help")

    assert result.returncode == 0
    assert "input" in result.stdout.lower()
    for option in (
        "-o",
        "--out",
        "--stdout",
        "--no-images",
        "--no-ocr",
        "--password",
        "--verify",
        "-q",
        "--quiet",
        "-v",
        "--version",
        "-h",
        "--help",
    ):
        assert option in result.stdout


def test_missing_input() -> None:
    result = run_docmd()

    assert result.returncode != 0
    assert "usage:" in result.stderr.lower()


def test_output_classification(tmp_path: Path) -> None:
    document = Document()
    document.add_paragraph("Body text")
    source = tmp_path / "sample.docx"
    document.save(source)

    file_output = tmp_path / "report.md"
    file_result = run_docmd(str(source), "-o", str(file_output))

    assert file_result.returncode == 0, file_result.stderr
    assert file_output.is_file()
    assert "Body text" in file_output.read_text(encoding="utf-8")
    assert not (tmp_path / "sample.md").exists()

    directory_output = tmp_path / "out"
    directory_result = run_docmd(str(source), "-o", str(directory_output))

    assert directory_result.returncode == 0, directory_result.stderr
    assert (directory_output / "sample.md").is_file()


def test_quiet_progress(tmp_path: Path) -> None:
    document = Document()
    document.add_paragraph("Body text")
    source = tmp_path / "input" / "sample.docx"
    source.parent.mkdir(parents=True)
    document.save(source)
    out = tmp_path / "out"

    result = run_docmd(str(source.parent), "-o", str(out), "--no-ocr")

    assert result.returncode == 0, result.stderr
    assert "sample.md" in result.stderr
    assert str(out) in result.stderr

    quiet_result = run_docmd(str(source.parent), "-o", str(out), "--no-ocr", "-q")

    assert quiet_result.returncode == 0, quiet_result.stderr
    assert quiet_result.stderr.strip() == ""

    broken = tmp_path / "broken.docx"
    broken.write_bytes(b"not a document")
    error_result = run_docmd(str(broken), "--stdout", "-q")

    assert error_result.returncode != 0
    assert "error" in error_result.stderr.lower()


def test_stdout_verify(tmp_path: Path) -> None:
    document = Document()
    document.add_paragraph("Body text")
    source = tmp_path / "ok.docx"
    document.save(source)

    ok = run_docmd(str(source), "--stdout", "--verify")

    assert ok.returncode == 0, ok.stderr
    assert "Body text" in ok.stdout

    broken = Document()
    broken.add_paragraph("```")
    broken_source = tmp_path / "broken.docx"
    broken.save(broken_source)

    bad = run_docmd(str(broken_source), "--stdout", "--verify")

    assert bad.returncode == 1
    assert "unbalanced code fence" in bad.stderr
    assert bad.stdout == ""
    assert not (tmp_path / "broken.md").exists()
    assert not (tmp_path / "broken.assets").exists()
